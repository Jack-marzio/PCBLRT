# -*- coding: utf-8 -*-
"""
report_creator.py
Creazione di report PDF e XLSX delle comunicazioni radio.
"""

import os
import re
import textwrap
from datetime import datetime

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font


class ReportCreator:
    """
    Crea report PDF e XLSX a partire dal log delle comunicazioni radio.
    Fornisce anche metodi di utilità per contare file esistenti.
    """

    # ------------------------------------------------------------------
    # Static / class helpers
    # ------------------------------------------------------------------

    @staticmethod
    def count_files(directory: str, extension: str) -> int:
        """Conta i file con una data estensione nella directory e restituisce il prossimo indice."""
        try:
            files = [f for f in os.listdir(directory) if f.endswith(extension)]
            return len(files) + 1
        except FileNotFoundError:
            print(f"Directory '{directory}' non trovata.")
            return 1

    @staticmethod
    def sanitize_filename(name: str) -> str:
        """Rimuove i caratteri non validi dal nome del file."""
        return re.sub(r'[<>:"/\\|?*]', "", name)

    @staticmethod
    def make_directory(directory: str) -> str | None:
        """Crea la directory se non esiste. Restituisce il percorso o None in caso di errore."""
        if not os.path.exists(directory):
            try:
                os.makedirs(directory)
                print(f"Directory creata: {directory}")
            except OSError as e:
                print(f"Errore creazione directory {directory}: {e}")
                return None
        return directory

    # ------------------------------------------------------------------
    # PDF
    # ------------------------------------------------------------------

    def create_pdf(self, operator: str, event: str, login_datetime: str,
                   communication_log: list, directory: str,
                   n_files: int, logo_path: str) -> str:
        """
        Genera il report PDF e lo salva nella directory indicata.
        Restituisce il percorso del file creato.
        """

        class _PDF(FPDF):
            COL_WIDTHS  = [35, 30, 30, 90, 90]
            HEADERS     = ["Data/Time", "Mittente", "Ricevente",
                           "Messaggio ricevuto", "Messaggio trasmesso"]

            def header(self):
                if hasattr(self, "logo_path") and os.path.exists(self.logo_path):
                    self.image(self.logo_path, 10, 8, 25)
                self.set_font("Arial", "B", 16)
                self.set_xy(0, 15)
                self.cell(0, 10, self.title_report, border=0, ln=True, align="C")
                self.ln(10)
                self.set_font("Arial", "", 10)
                self.cell(0, 6, f"Operatore: {self.operator}", ln=True)
                self.cell(0, 6, f"Evento: {self.event}", ln=True)
                self.ln(5)
                self._draw_table_header()

            def footer(self):
                self.set_y(-12)
                self.set_font("Arial", "I", 8)
                self.cell(0, 5, f"Pagina {self.page_no()}", align="C")

            def _draw_table_header(self):
                self.set_font("Arial", "B", 9)
                self.set_fill_color(230, 230, 230)
                for w, h in zip(self.COL_WIDTHS, self.HEADERS):
                    self.cell(w, 8, h, border=1, align="C", fill=True)
                self.ln()

        pdf = _PDF(orientation="L")
        pdf.logo_path    = logo_path
        pdf.operator     = operator
        pdf.event        = event
        pdf.title_report = (
            f"Rapporto registrazione comunicazioni radio "
            f"n° {n_files}/{login_datetime[:4]}"
        )
        pdf.add_page()
        pdf.set_font("Arial", "", 9)

        line_height = 5
        col_widths  = _PDF.COL_WIDTHS

        for entry in communication_log:
            row = ["\n".join(textwrap.wrap(str(x), 40)) for x in entry]
            max_lines  = max(r.count("\n") + 1 for r in row)
            row_height = line_height * max_lines

            if pdf.get_y() + row_height > pdf.page_break_trigger:
                pdf.add_page()

            x_start = pdf.get_x()
            y_start = pdf.get_y()
            for i, (cell_text, w) in enumerate(zip(row, col_widths)):
                pdf.set_xy(x_start, y_start)
                pdf.multi_cell(w, line_height, cell_text,
                               border=1, align="C" if i < 3 else "L")
                x_start += w
            pdf.set_y(y_start + row_height)

        file_name = (
            f"rapporto_comunicazioni_radio_"
            f"{self.sanitize_filename(operator)}_{n_files}_{login_datetime[:4]}.pdf"
        )
        file_path = os.path.join(directory, file_name)
        try:
            pdf.output(file_path)
            print(f"PDF salvato: {file_path}")
        except Exception as e:
            print(f"Errore salvataggio PDF: {e}")
        return file_path

    # ------------------------------------------------------------------
    # XLSX
    # ------------------------------------------------------------------

    def create_xlsx(self, operator: str, event: str, login_datetime: str,
                    communication_log: list, directory: str,
                    n_files: int, logo_path: str) -> str:
        """
        Genera il report XLSX e lo salva nella directory indicata.
        Restituisce il percorso del file creato.
        """
        IMG_W, IMG_H = 150, 150
        TEXT_SIZE    = 12
        TABLE_HEADER = ["Data/Ora", "Mittente", "Ricevente",
                        "Messaggio ricevuto", "Messaggio trasmesso"]

        wb = Workbook()
        ws = wb.active

        img = XLImage(logo_path)
        img.width, img.height = IMG_W, IMG_H
        ws.add_image(img, "A3")

        ws.append([operator, event, login_datetime, n_files])
        ws.append(TABLE_HEADER)

        for col in range(1, len(TABLE_HEADER) + 1):
            cell = ws.cell(row=3, column=col)
            cell.font      = Font(bold=True, size=TEXT_SIZE)
            cell.alignment = Alignment(horizontal="center")

        max_lengths = [len(h) for h in TABLE_HEADER]
        for entry in communication_log:
            row_data = list(entry)
            for i, value in enumerate(row_data):
                max_lengths[i] = max(max_lengths[i], len(str(value)))
            ws.append(row_data)

        max_width = 20
        for i, length in enumerate(max_lengths, start=1):
            col_letter = ws.cell(row=1, column=i).column_letter
            if length > max_width:
                ws.column_dimensions[col_letter].width = max_width
                for row in ws.iter_rows(min_col=i, max_col=i):
                    for cell in row:
                        cell.alignment = Alignment(wrap_text=True)
            else:
                ws.column_dimensions[col_letter].width = length + 2

        file_name = (
            f"rapporto_comunicazioni_radio_"
            f"{self.sanitize_filename(operator)}_{n_files}_{login_datetime[:4]}.xlsx"
        )
        file_path = os.path.join(directory, file_name)
        try:
            wb.save(file_path)
            print(f"XLSX salvato: {file_path}")
        except Exception as e:
            print(f"Errore salvataggio XLSX: {e}")
        return file_path
